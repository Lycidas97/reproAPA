#!/usr/bin/env python3
"""Independent read-back of Supplementary_Tables_S1-S8.xlsx.

Checks: sheet order; every block's header and every cell against its source file (after the documented row
filter); Excel Table ranges in the package XML; fonts in styles.xml; key manuscript numbers.
"""
import csv, json, re, sys, zipfile
from pathlib import Path
import openpyxl, pandas as pd

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import build_workbooks as B  # noqa: E402
from build_supplement_xlsx import OUTNAME  # noqa: E402
import amendments_1530 as A  # noqa: E402
import amendments_1008 as A2  # noqa: E402

X = HERE / OUTNAME
plan = list(csv.DictReader(open(HERE / 'table-plan.tsv'), delimiter='\t'))
blocks = list(csv.DictReader(open(HERE / 'supplement-blocks.tsv'), delimiter='\t'))
src = {(r['sheet'], r['path']): r for r in csv.DictReader(open(HERE / 'table-sources.tsv'), delimiter='\t')
       if r['layer'] == 'sheet'}
problems, maxrel, ncell = [], 0.0, 0
BLK = {}
ROWS = {}

wb = openpyxl.load_workbook(X, read_only=True)
exp = ['Contents'] + [s['sheet'] for s in plan] + ['Sources']
if wb.sheetnames != exp:
    problems.append('sheet order differs')
by_sheet = {}
for b in blocks:
    by_sheet.setdefault(b['sheet'], []).append(b)
for sh, bl in by_sheet.items():
    rows = list(wb[sh].iter_rows(values_only=True))
    ROWS[sh] = rows
    for b in bl:
        if b['role'] == 'derived':
            d = dict((n, f) for n, f, _, _ in A.DERIVED[sh])[b['path']]()
        else:
            r = src.get((sh, b['path']), {'path': b['path'], 'role': 'core_list'})
            d = pd.read_csv(B.src_path(r), sep='\t', dtype=str, keep_default_na=False, na_filter=False,
                            compression='gzip' if b['path'].endswith('.gz') else None)
            flt = A.ADD.get((sh, b['path'])) or A.FILTER.get((sh, b['path'])) or B.FILTERS.get(b['path'])
            if flt:
                d = flt[1](d)
            for cname, fn, _ in A.ADD_COLS.get((sh, b['path']), []):
                d = d.copy(); d[cname] = fn(d)
            if (sh, b['path']) in A2.DROP_COLS:
                d = d.drop(columns=A2.DROP_COLS[(sh, b['path'])][0])
        BLK[b['block']] = (h0 := int(b['header_row']) - 1, len(d), list(d.columns), sh)
        h = int(b['header_row']) - 1
        nc = len(d.columns)
        if [rows[h][j] for j in range(nc)] != list(d.columns):
            problems.append(f"{b['block']}: header")
        body = rows[h + 1:h + 1 + len(d)]
        if len(body) != len(d):
            problems.append(f"{b['block']}: rows {len(body)} vs {len(d)}")
        mism = 0
        for a, v in zip(body, d.itertuples(index=False)):
            for j in range(nc):
                u, w = a[j] if j < len(a) else None, v[j]
                ncell += 1
                if w == '':
                    ok = u is None
                elif isinstance(u, (int, float)) and not isinstance(u, bool):
                    fw = float(w)
                    rel = abs(fw - u) / max(abs(fw), 1e-300)
                    maxrel = max(maxrel, rel)
                    ok = rel <= 1e-15
                else:
                    ok = str(u) == w
                mism += not ok
        # the row after the block must be empty (no spill-over into the next block)
        after = rows[h + 1 + len(d)] if h + 1 + len(d) < len(rows) else ()
        if any(x is not None for x in after):
            problems.append(f"{b['block']}: row after block not empty")
        if mism:
            problems.append(f"{b['block']}: {mism} cells differ")
wb.close()

# Excel Table objects: ranges must match header row, row count and width
z = zipfile.ZipFile(X)
refs = {}
for n in z.namelist():
    if n.startswith('xl/tables/table'):
        t = z.read(n).decode()
        refs[re.search(r'name="([^"]+)"', t).group(1)] = re.search(r'ref="([^"]+)"', t).group(1)
fonts = set(re.findall(r'<name val="([^"]+)"', z.read('xl/styles.xml').decode()))
from openpyxl.utils import get_column_letter
n_tab = 0
for b in blocks:
    if b['excel_table'] != 'True':
        continue
    n_tab += 1
    i = b['block'].split('.')[-1] if '.' in b['block'] else '1'
    name = f"T_{b['sheet']}_{i}"
    h = int(b['header_row'])
    want = f"A{h}:{get_column_letter(int(b['columns']))}{h + max(int(b['rows']), 1)}"
    if refs.get(name) != want:
        problems.append(f'table range {name} {refs.get(name)} != {want}')

out = {'blocks': len(blocks), 'cells_checked': ncell, 'max_relative_numeric_difference': maxrel,
       'excel_tables': len(refs), 'excel_tables_expected': n_tab + 1, 'fonts_in_styles': sorted(fonts),
       'plain_blocks': [b['block'] for b in blocks if b['excel_table'] != 'True'], 'problems': problems}
json.dump(out, open(HERE / 'supplement-verify.json', 'w'), indent=1)
print(json.dumps(out, indent=1))

# ---- manuscript numbers, read back from the workbook itself (15:30 check items)
def block(label):
    h, n, cols, sh = BLK[label]
    rows = ROWS[sh]
    return pd.DataFrame([list(r[:len(cols)]) for r in rows[h + 1:h + 1 + n]], columns=cols)


def find(sh, path):
    return next(b['block'] for b in blocks if b['sheet'] == sh and b['path'] == path)


M = {}
c = block(find('S4c', 'intermediate_state_class_counts'))
fig = json.load(open(A.ROOT / 'notes/2026-10-06-supplement-polyapipe/candidates/D13/SuppFig.S16.values.json'))['c']
key = {('human', 'R'): 'Human R', ('mouse', 'R'): 'Mouse R', ('mouse', 'E'): 'Mouse E'}
M['S4c counts == Figure S16c values'] = all(
    {k: int(v) for k, v in row.items() if k not in ('species', 'intermediate_state')} == fig[key[(row['species'], row['intermediate_state'])]]
    for row in c.to_dict('records'))
M['S4c aligned / reverse / denominators'] = [(int(r['endpoint_aligned']), int(r['reverse']), int(r['total'])) for r in c.to_dict('records')]
g = block(find('S4c', 'intermediate/process/same_series_process_classes.tsv.gz'))
M['S4c gene rows agree with counts'] = g.groupby(['species', 'state']).size().to_dict() == {('human', 'R'): 483, ('mouse', 'E'): 604, ('mouse', 'R'): 604}
a = block(find('S5a', 'direction_difference/direction_contrast_tests.tsv'))
b2 = block(find('S5a', 'direction_difference/original_pair_direction_sensitivity.tsv'))
M['S5a direction-difference family size (48 + 32)'] = (len(a), len(b2), len(a) + len(b2))
M['S5a min q_complete_80'] = min(min(a['q_complete_80']), min(b2['q_complete_80']))
e = block(find('S5d', 'same_pas/v2_primary_same_axis_gene_effects.tsv.gz'))
M['S5d mass_min by pair'] = {f"{r['design']} {r['pair']}": round(r['mass_min'], 6) for r in e.to_dict('records')}
u = block(find('S5d', 'same_pas/v2_primary_annotated_usage.tsv.gz'))
w = u[(u.code == 'W') & (u.design == 'PMS')].sort_values('S', ascending=False).head(2)
h1 = u[(u.code == 'H1') & (u.design == 'PMS')].sort_values('S', ascending=False).head(2)
M['S5d two-site usage, H1_W pair (S, P, M)'] = {sp: [round(x[c].sum(), 4) for c in ('S', 'P', 'M')] for sp, x in (('human H1', h1), ('mouse W', w))}
iso = block(find('S5d', 'alternative_terminal_exon_isoforms'))
M['S5d isoforms (aa)'] = dict(zip(iso.transcript_id, iso.protein_aa))
sep = block(find('S5d', 'matched_PAS_separation'))
M['S5d separation (nt)'] = dict(zip(sep.species, sep.separation_nt))
M['S1a marker blocks (rows)'] = [int(b['rows']) for b in blocks if b['sheet'] == 'S1a' and 'marker_by_state' in b['path']]
M['S4a design block rows'] = int(next(b['rows'] for b in blocks if b['sheet'] == 'S4a' and b['path'] == 'global/planned_route_contrasts.tsv'))
t = block(find('S7b', 'examples/inclusion_effect_sensitivity.tsv'))
M['S7b inclusion genes'] = sorted(set(t.gene))
f2 = block(find('S3b', 'held-out_route_mean_correlations'))
M['S3b r range by comparison'] = {k: (round(x.r.min(), 3), round(x.r.max(), 3), len(x)) for k, x in f2.groupby('held_out_route_means_compared_with')}
rs = block(find('S6c', 'weighted_retention/retention_summary.tsv'))
M['S6c mouse RBP genes without complete sites'] = rs[(rs.species == 'mouse') & (rs.category == 'RBP')].genes_without_complete_sites.tolist()
wbt = openpyxl.load_workbook(X, read_only=True)
M['titles S2d / S4a'] = [wbt[x]['A1'].value for x in ('S2d', 'S4a')]
M['2026-10-08: excluded blocks absent'] = not any((b['sheet'], b['path']) in A2.EXCLUDE for b in blocks)
M['2026-10-08: S6b blocks'] = [(b['block'], b['path']) for b in blocks if b['sheet'] == 'S6b']
M['2026-10-08: S3e.1 has no H column'] = not any('hidden_remodeling' in c for c in BLK[find('S3e', 'profiles/architecture_route_summary.tsv')][2])
M['2026-10-08: no inference/primary_test columns in S6b'] = not any(c in ('status', 'primary_test', 'gene_label_q_all_rows') for b in blocks if b['sheet'] == 'S6b' for c in BLK[b['block']][2])
M['2026-10-08: Contents A3 is the contrast key'] = wbt['Contents']['A3'].value == A2.CONTRAST_KEY
M['2026-10-08: Sources A2'] = wbt['Sources']['A2'].value == A2.SOURCES_NOTE
s7 = BLK[find('S7b', 'trip12_evidence/figure5/site-display-evidence.tsv')]
M['2026-10-08: S7b displayed_in_F5F note'] = ROWS['S7b'][s7[0] - 1][0] == A2.NOTES[('S7b', 'trip12_evidence/figure5/site-display-evidence.tsv')][0]
M['2026-10-08: S2d related figures'] = wbt['S2d']['A2'].value
out['manuscript_checks'] = M
json.dump(out, open(HERE / 'supplement-verify.json', 'w'), indent=1, default=str)
print(json.dumps(M, indent=1, default=str))
sys.exit(1 if problems else 0)
