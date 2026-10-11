"""Extract published support flags and passages; verify the GEO-based audit.

Uses openpyxl from the bundled Codex Python runtime. The published spreadsheet
is corroborating evidence: its four extra IDs lack structures in the GEO GTF
and are not substituted into the primary structural analysis.
"""
from pathlib import Path
from collections import Counter, defaultdict
import csv
import json
import re
import openpyxl
from analyze import HERE, write


def rows(name):
    return list(csv.DictReader((HERE / name).open(), delimiter='\t'))


def normalized(name):
    return re.sub(r'_(pcs_r10|r9|pc10)$', '', name.lower())


def main():
    w = openpyxl.load_workbook(HERE / 'raw/supplement-data1.xlsx', read_only=True, data_only=True)
    for name in ['lrRNAseq-samples', 'Statistics-lrRNAseq-reads']:
        with (HERE / (name + '.tsv')).open('w') as f:
            csv.writer(f, delimiter='\t').writerows(w[name].iter_rows(values_only=True))
    w.close()
    w = openpyxl.load_workbook(HERE / 'raw/supplement-data2.xlsx', read_only=True, data_only=True)
    iterator = w['Classification-file'].iter_rows(values_only=True)
    header = next(iterator)
    with (HERE / 'filtered_supplement_classification.tsv').open('w') as f:
        writer = csv.writer(f, delimiter='\t')
        writer.writerow(header)
        for row in iterator:
            if str(row[1]).startswith(('Rad23b.', 'Hmgn2.', 'Eif4ebp2.')):
                writer.writerow(row)
    w.close()
    pub = {r['Transcript name ID']: r for r in rows('filtered_supplement_classification.tsv')}
    audit = {r['transcript_id']: r for r in rows('transcript_audit.tsv')}
    support = []
    for tid, r in audit.items():
        q = pub[tid]
        assert q['Filtered out for low expression'] == 'False'
        assert q['Supported either by junction or TSS'] == 'True'
        assert q['Structural category'] == r['structural_category']
        assert int(q['Number of exons']) == int(r['n_exons'])
        assert q['All canonical'] == r['all_canonical']
        assert q['polyA motif'] == r['polyA_motif']
        support.append(dict(gene=r['gene'], transcript_id=tid, published_retained=True,
                            junctions_supported=q['All junctions supported merge'], TSS_supported=q['TSS supported'],
                            predicted_coding=q['Coding'], subcategory=q['Structural subcategory']))
    write('published_candidate_support.tsv', support)
    passages = {(normalized(r['Name of the lrRNAseq sample']), r['Flow cell generation']): r
                for r in rows('lrRNAseq-samples.tsv')}
    inventory = rows('sample_inventory.tsv')
    enriched = []
    for r in inventory:
        q = passages[normalized(r['TPM_column']), r['platform']]
        enriched.append(dict(**r, timepoint=q['Timepoint'], passage=q['Passage'],
                             published_cell_line=q['Cell line'], published_medium=q['Medium']))
    assert len(enriched) == len(passages) == 45
    write('sample_inventory_with_passage.tsv', enriched)
    groups = defaultdict(list)
    for r in enriched:
        groups[r['published_cell_line'], r['timepoint'], r['published_medium']].append(r)
    write('sample_design_groups.tsv', [dict(cell_line=k[0], timepoint=k[1], medium=k[2],
                                           n_GSM=len(v), platforms=';'.join(r['platform'] for r in v),
                                           passages=';'.join(r['passage'] for r in v),
                                           GSMs=';'.join(r['GSM'] for r in v)) for k, v in groups.items()])
    tpm = {r['transcript_id']: r for r in rows('candidate_TPM.tsv')}
    summary = []
    for r in rows('compatible_isoform_pairs.tsv'):
        a, b = tpm[r['included_transcript']], tpm[r['proximal_transcript']]
        for platform in ['all', 'R9', 'R10']:
            selected = [s for s in enriched if platform == 'all' or s['platform'] == platform]
            n = sum(float(a[s['TPM_column']]) >= 1 and float(b[s['TPM_column']]) >= 1 for s in selected)
            r['both_TPM_ge1_samples_' + platform] = n
        summary.append(r)
    write('compatible_isoform_pairs_supported.tsv', summary)
    verification = dict(status='PASS', all_76_GEO_candidates_retained_and_supported_in_published_table=True,
                        all_45_TPM_samples_join_published_passage_table=True,
                        published_table_only_ids=sorted(set(pub) - set(audit)),
                        published_table_only_ids_used_for_structure=False,
                        low_expression_rule='TPM >=1 in at least 3 samples; published final support flags checked',
                        interpretation='Different passages supply culture context, not a verified replicate hierarchy')
    (HERE / 'supplement_verification.json').write_text(json.dumps(verification, indent=2) + '\n')
    print(json.dumps(verification, indent=2), flush=True)


if __name__ == '__main__':
    main()
