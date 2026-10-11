import openpyxl, zipfile, re, json, pandas as pd, math, hashlib
import os
ROOT=os.environ.get('REPRO_ROOT','/mnt/user-data/uploads/reproAPA'); OD=os.environ.get('S9F_OUT','/mnt/user-data/outputs/s9f')
A=ROOT+'/table/Supplementary_Tables_S1-S9.xlsx'; B=OD+'/Supplementary_Tables_S1-S9.xlsx'
wa=openpyxl.load_workbook(A,read_only=True); wb=openpyxl.load_workbook(B,read_only=True)
res={'sheets_original':wa.sheetnames,'sheets_new':wb.sheetnames,'diff':{}}
for n in wa.sheetnames:
    a=wa[n]; b=wb[n]; nd=0; nc=0
    for ra,rb in __import__('itertools').zip_longest(a.iter_rows(values_only=True),b.iter_rows(values_only=True),fillvalue=()):
        ra=list(ra); rb=list(rb); L=max(len(ra),len(rb)); ra+= [None]*(L-len(ra)); rb+=[None]*(L-len(rb))
        for x,y in zip(ra,rb):
            nc+=1
            if x!=y and not (isinstance(x,float) and isinstance(y,float) and (x==y or (math.isnan(x) and math.isnan(y)))): nd+=1
    res['diff'][n]=dict(cells=nc,different=nd)
def tables(p):
    z=zipfile.ZipFile(p); return sorted(re.search(r'displayName="([^"]+)"',z.read(f).decode()).group(1) for f in z.namelist() if re.match(r'xl/tables/table\d+\.xml$',f))
ta,tb=tables(A),tables(B); res['tables_original']=len(ta); res['tables_new']=len(tb); res['tables_missing']=sorted(set(ta)-set(tb)); res['tables_added']=sorted(set(tb)-set(ta))
# S9f values against source TSVs
D=ROOT+'/notes/2026-10-10-marker-reference-sensitivity/'
ws=wb['S9f']; rows=list(ws.iter_rows(values_only=True))
blocks=[i for i,r in enumerate(rows) if r and isinstance(r[0],str) and r[0].startswith('S9f.')]
chk=[]
specs=[('unique_support_bridge_route_summary.tsv',"background == 'common_M_bridge' and version == 'core_third_strict' and setting == 'supported50' and threshold == 0.02"),
       ('unique_support_bridge_route_summary.tsv',"background == 'common_M_bridge'"),('unique_support_bridge_route_summary.tsv',"background == 'common_P_bridge'"),
       ('unique_support_route_summary.tsv',None),('gradient_route_summary.tsv',None),('classification_confusion.tsv',None)]
for bi,(f,q) in zip(blocks,specs):
    df=pd.read_csv(D+f,sep='\t'); df=df.query(q) if q else df
    hdr=list(rows[bi+3][:len(df.columns)]); assert hdr==list(df.columns),(f,hdr[:3])
    bad=0
    for k,(_,r) in enumerate(df.iterrows()):
        got=rows[bi+4+k][:len(df.columns)]
        for x,y in zip(r.values,got):
            if isinstance(x,float) and math.isnan(x): ok= y=='NA'
            elif isinstance(x,(bool,)) or type(x).__name__=='bool_': ok= y==str(bool(x))
            elif isinstance(x,(int,float)) or hasattr(x,'dtype'): ok= y is not None and not isinstance(y,str) and float(x)==float(y)
            else: ok= str(x)==y
            bad+= not ok
    chk.append(dict(block=rows[bi][0][:6],rows=len(df),bad_cells=bad))
res['S9f_blocks']=chk
c=wb['Contents']; res['contents_last']=[r for r in c.iter_rows(values_only=True) if r[0]][-2:]
s=wb['Sources']; res['sources_last']=[str(r[:4]) for r in s.iter_rows(values_only=True)][-6:]
res['sha_original']=hashlib.sha256(open(A,'rb').read()).hexdigest(); res['sha_new']=hashlib.sha256(open(B,'rb').read()).hexdigest()
json.dump(res,open(OD+'/verify.json','w'),indent=1,default=str)
print(json.dumps({k:v for k,v in res.items() if k!='diff'},indent=1,default=str)[:2500]); print('sheets with differences:',{k:v for k,v in res['diff'].items() if v['different']})
za=zipfile.ZipFile(A); zb=zipfile.ZipFile(B)
changed=[n for n in za.namelist() if za.read(n)!=zb.read(n)]; added=[n for n in zb.namelist() if n not in za.namelist()]
res['zip_parts_changed']=changed; res['zip_parts_added']=added
json.dump(res,open(OD+'/verify.json','w'),indent=1,default=str)
print('changed parts',changed); print('added',added)
